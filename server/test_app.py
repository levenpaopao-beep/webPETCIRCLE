"""隔离数据库的业务与权限回归；不改正式或本机经营资料。"""
import io,json,secrets,shutil,tempfile,threading,unittest,urllib.request,urllib.error,http.cookiejar,zipfile
from pathlib import Path
from http.server import ThreadingHTTPServer
from datetime import datetime,timezone,timedelta
from server import app

class SiteTests(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.tmp=tempfile.TemporaryDirectory();cls.old=(app.STATE,app.DB);app.STATE=Path(cls.tmp.name);app.DB=app.STATE/'test.sqlite';app.init()
  cls.config=json.loads((app.ROOT/'.local/source-settings.json').read_text());cls.settings={k:cls.config[k] for k in ['cn','fob']}
  cls.original=(app.ROOT/'.local/source-snapshot.json').read_text()
  with app.connect() as c:
   for k,v in [('settings',json.dumps(cls.settings)),('source_time',datetime.now(timezone.utc).isoformat())]:c.execute('INSERT INTO metadata VALUES(?,?)',(k,v))
   for code,price in [('A',14),('B',20),('C',11.03)]:
    c.execute('INSERT INTO products(code,name_zh,name_en,source_category,category,published,photo) VALUES(?,?,?,?,?,1,?)',(code,'测试毛衣'+code,'Knit sweater '+code,'毛衣','knitwear',code+'.jpg'))
    c.execute('INSERT INTO variants VALUES(?,?,?,?,?,?,?,?,?)',(code+'-S',code,'红色','S',price,10000,120,50,'40×30×30 cm'))
    c.execute('INSERT INTO variants VALUES(?,?,?,?,?,?,?,?,?)',(code+'-M',code,'红色','M',price,200,130,None,''))
    shutil.copy2(next((app.ROOT/'.local/images').glob('*.jpg')),app.STATE/'images'/f'{code}.jpg')
  cls.server=ThreadingHTTPServer(('127.0.0.1',0),app.Handler);cls.base='http://127.0.0.1:'+str(cls.server.server_port)
  cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True);cls.thread.start()
  cls.admin=cls.client();password=(app.STATE/'管理员登录.txt').read_text().split('密码：')[1].strip()
  cls.req(cls.admin,'/login',{'email':'admin@petcircle.local','password':password})
 @classmethod
 def tearDownClass(cls):
  cls.server.shutdown();app.STATE,app.DB=cls.old;cls.tmp.cleanup()
 @staticmethod
 def client():return urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
 @classmethod
 def req(cls,client,path,body=None,status=200,raw=False,origin=None):
  headers={'Content-Type':'application/json'}
  if origin:headers['Origin']=origin
  r=urllib.request.Request(cls.base+'/api'+path,data=json.dumps(body).encode() if body is not None else None,headers=headers)
  try:res=client.open(r)
  except urllib.error.HTTPError as e:res=e
  assert res.status==status,(path,res.status,res.read())
  return res.read() if raw else json.loads(res.read())
 def customer(self):
  c=self.client();d=self.req(c,'/register',{'email':secrets.token_hex(5)+'@example.com','password':'safepassword12','contact':'Buyer','company':'Sample shop','country':'United Kingdom'},201 if False else 200);return c,d['user']
 def create(self,c,lines=None,**extra):return self.req(c,'/quotes',{'lines':lines or [{'sku':'A-S','quantity':2000}],'requestKey':secrets.token_hex(10),'language':'en',**extra},201)['quote']
 def test_guest_message_lifecycle(self):
  app.MESSAGE_ATTEMPTS.clear();guest=self.client();key=secrets.token_hex(10)
  body={'name':'留言测试','email':'message@example.com','phone':'','wechat':'','company':'测试公司','message':'希望讨论贴牌\n与样品','language':'en','source':'oem','requestKey':key}
  self.req(guest,'/messages',body,201)
  self.req(guest,'/messages',body)
  with app.connect() as db:self.assertEqual(db.execute('SELECT count(*) FROM messages WHERE request_key=?',(key,)).fetchone()[0],1)
  self.req(guest,'/admin/messages',status=403)
  customer,_=self.customer();self.req(customer,'/admin/messages',status=403)
  rows=self.req(self.admin,'/admin/messages')['messages'];msg=next(m for m in rows if m['name']=='留言测试')
  self.assertEqual(msg['message'],body['message']);self.assertEqual(msg['unread'],1);self.assertNotIn('request_key',msg)
  self.assertGreater(self.req(self.admin,'/admin/summary')['unreadMessages'],0)
  self.req(customer,'/admin/messages',{'id':msg['id']},403)
  self.req(self.admin,'/admin/messages',{'id':msg['id']})
  self.assertEqual(next(m for m in self.req(self.admin,'/admin/messages')['messages'] if m['id']==msg['id'])['unread'],0)
  # SQLite重新连接仍可读到完整内容。
  with app.connect() as db:self.assertEqual(db.execute('SELECT email FROM messages WHERE id=?',(msg['id'],)).fetchone()[0],body['email'])
 def test_message_validation_and_contact_methods(self):
  app.MESSAGE_ATTEMPTS.clear();guest=self.client()
  base={'name':'Buyer','requestKey':secrets.token_hex(10)}
  for fields in [{},{'name':'' ,'email':'a@example.com'},{'email':'bad'},{'phone':'abc'},{'message':'x'*3001}]:
   self.req(guest,'/messages',{**base,**fields},400)
  for method,value in [('email','valid@example.com'),('phone','+852 12345678'),('wechat','wechat-test')]:
   self.req(guest,'/messages',{**base,method:value,'requestKey':secrets.token_hex(10)},201)
  self.req(guest,'/messages',{**base,'email':'valid@example.com'},403,origin='https://evil.example')
 def test_message_rate_limit(self):
  app.MESSAGE_ATTEMPTS.clear();guest=self.client()
  try:
   for i in range(10):self.req(guest,'/messages',{'name':'Rate test','wechat':'contact','requestKey':secrets.token_hex(10)},201)
   self.req(guest,'/messages',{'name':'Rate test','wechat':'contact','requestKey':secrets.token_hex(10)},429)
  finally:app.MESSAGE_ATTEMPTS.clear()
 def test_public_privacy(self):
  guest=self.client();data=self.req(guest,'/products');text=json.dumps(data)
  for forbidden in ['exchangeRate','markupRate','available','unitWeightG','password','account']:self.assertNotIn('"'+forbidden+'"',text)
  self.assertTrue(data['products']);self.assertEqual(data['products'][0]['currency'],'USD')
  self.req(guest,'/products?admin=1',status=403);self.req(guest,'/admin/customers',status=403)
 def test_registration_and_authorization(self):
  c,u=self.customer();self.assertEqual(self.req(c,'/me')['user']['email'],u['email']);self.req(c,'/admin/quotes',status=403)
  self.req(self.client(),'/register',{'email':'bad','password':'short'},400)
  self.req(c,'/profile',{'contact':'Buyer2','company':'Shop2','country':'US','phone':''});self.assertEqual(self.req(c,'/me')['user']['contact'],'Buyer2')
  self.req(c,'/logout',{});self.req(c,'/quotes',status=403)
 def test_origin_guard(self):self.req(self.client(),'/register',{'email':'x@example.com','password':'safepassword12'},403,origin='https://evil.example')
 def test_minimums_and_integer_stock(self):
  c,u=self.customer()
  for lines in [[{'sku':'A-S','quantity':99},{'sku':'B-S','quantity':2000}],[{'sku':'A-S','quantity':1999}],[{'sku':'A-S','quantity':2.5}],[{'sku':'A-S','quantity':-2000}]]:
   self.req(c,'/quotes',{'lines':lines,'requestKey':secrets.token_hex(10)},400)
  with app.connect() as db:db.execute("UPDATE variants SET available=1 WHERE sku='A-M'")
  try:
   d=self.req(self.client(),'/stock-check',{'lines':[{'sku':'A-M','quantity':2000}]});self.assertTrue(d['results'][0]['ok']);self.assertNotIn('available',json.dumps(d))
   self.assertEqual(self.create(c,[{'sku':'A-M','quantity':2000}])['quantity'],2000)
   for stock in [0,None,-1]:
    with app.connect() as db:db.execute("UPDATE variants SET available=? WHERE sku='A-M'",(stock,))
    self.assertFalse(self.req(self.client(),'/stock-check',{'lines':[{'sku':'A-M','quantity':2000}]})['results'][0]['ok'])
    self.req(c,'/quotes',{'lines':[{'sku':'A-M','quantity':2000}],'requestKey':secrets.token_hex(10)},400)
  finally:
   with app.connect() as db:db.execute("UPDATE variants SET available=200 WHERE sku='A-M'")
 def test_tiers_and_duplicate_aggregation(self):
  c,u=self.customer()
  for qty,rate in [(299,1),(300,.97),(500,.95)]:
   q=self.create(c,[{'sku':'A-S','quantity':qty},{'sku':'B-S','quantity':2000}]);self.assertEqual(q['lines'][0]['unitPrice'],app.money(14/7*1.08*rate))
  q=self.create(c,[{'sku':'A-S','quantity':1000},{'sku':'A-S','quantity':1000}]);self.assertEqual(len(q['lines']),1);self.assertEqual(q['quantity'],2000)
  self.assertEqual(q['total'],sum(x['amount'] for x in q['lines']))
 def test_first_unread_and_idempotency(self):
  c,u=self.customer();key=secrets.token_hex(10);q=self.create(c,requestKey=key);same=self.req(c,'/quotes',{'lines':[{'sku':'A-S','quantity':2000}],'requestKey':key})['quote'];self.assertEqual(q['id'],same['id']);self.assertTrue(q['firstInquiry']);self.assertTrue(q['unread'])
  other=self.create(c);self.assertFalse(other['firstInquiry']);seen=self.req(self.admin,'/quotes/'+str(q['id']))['quote'];self.assertFalse(seen['unread']);self.assertTrue(seen['firstInquiry'])
  self.req(self.admin,'/admin/quotes',{'id':q['id'],'status':'following'});self.assertEqual(self.req(c,'/quotes/'+str(q['id']))['quote']['status'],'following')
 def test_currency_thresholds(self):
  c,u=self.customer();self.req(self.admin,'/admin/customers',{'id':u['id'],'currency':'CNY'});self.assertEqual(self.req(c,'/me')['user']['currency'],'CNY')
  self.assertEqual(self.req(c,'/products')['products'][0]['currency'],'CNY')
  for qty,rate in [(100,1),(715,.95),(2143,.9)]:
   q=self.create(c,[{'sku':'A-S','quantity':qty}]);self.assertEqual(q['lines'][0]['unitPrice'],app.money(14*rate));self.assertEqual(q['currency'],'CNY')
  self.req(c,'/quotes',{'lines':[{'sku':'A-S','quantity':20}],'requestKey':secrets.token_hex(10)},400)
 def test_history_after_unpublish_and_other_customer(self):
  c,u=self.customer();q=self.create(c);other,_=self.customer();self.req(other,'/quotes/'+str(q['id']),status=404)
  try:
   self.req(self.admin,'/admin/products',{'codes':['A'],'changes':{'published':False}})
   self.req(c,'/products/A',status=404);history=self.req(c,'/quotes/'+str(q['id']))['quote'];self.assertEqual(history['total'],q['total']);self.assertTrue(self.req(c,'/quotes/'+str(q['id'])+'/image/A',raw=True).startswith(b'\xff\xd8'))
   self.req(c,'/quotes',{'lines':[{'sku':'A-S','quantity':2000}],'requestKey':secrets.token_hex(10)},400)
  finally:self.req(self.admin,'/admin/products',{'codes':['A'],'changes':{'published':True}})
 def test_stale_data_blocked(self):
  c,u=self.customer()
  try:
   with app.connect() as db:db.execute("UPDATE metadata SET value=? WHERE key='source_time'",((datetime.now(timezone.utc)-timedelta(days=3)).isoformat(),))
   self.req(c,'/quotes',{'lines':[{'sku':'A-S','quantity':2000}],'requestKey':secrets.token_hex(10)},400)
  finally:
   with app.connect() as db:db.execute("UPDATE metadata SET value=? WHERE key='source_time'",(datetime.now(timezone.utc).isoformat(),))
 def test_sync_preserves_lifecycle(self):
  customer,u=self.customer();quote=self.create(customer)
  with app.connect() as c:
   c.execute("UPDATE products SET featured=1,position=7,category='coats' WHERE code='A'")
   rows=[]
   for r in c.execute('SELECT v.*,p.name_zh,p.name_en FROM variants v JOIN products p ON p.code=v.code'):
    rows.append({'code':r['code'],'sku':r['sku'],'spec':r['color']+'/'+r['size'],'color':r['color'],'size':r['size'],'name_cn':r['name_zh'],'name_en':r['name_en'],'category':'毛衣','price':r['price'],'available':r['available'],'unit_weight_g':r['weight'],'carton_qty':r['carton_qty'],'carton_size':r['carton_size']})
  with app.connect() as c:
   c.execute("INSERT INTO products(code,name_zh,name_en,source_category,category,published) VALUES('MATERIAL','黄毛球','Yellow pompom','特殊辅料','accessories',1)")
  rows.append({'code':'MATERIAL','sku':'MATERIAL-S','category':'特殊辅料','name_cn':'黄毛球','available':100,'price':1})
  snapshot=app.STATE/'test-source.json' ;config=app.STATE/'test-settings.json';snapshot.write_text(json.dumps({'status':'success','completeness':'complete','synced_at':app.now(),'rows':rows}));config.write_text(json.dumps(self.settings))
  with app.connect() as c:
   if not c.execute("SELECT 1 FROM product_content WHERE code='B'").fetchone():app.save_product_content(c,{'code':'B','images':['B.jpg']})
   content_before=dict(c.execute("SELECT * FROM product_content WHERE code='B'").fetchone());photo_before=c.execute("SELECT photo FROM products WHERE code='B'").fetchone()[0]
  app.sync(snapshot,config)
  with app.connect() as c:
   self.assertEqual(dict(c.execute("SELECT * FROM product_content WHERE code='B'").fetchone()),content_before);self.assertEqual(c.execute("SELECT photo FROM products WHERE code='B'").fetchone()[0],photo_before)
   r=c.execute("SELECT * FROM products WHERE code='A'").fetchone();self.assertTrue(r['published']);self.assertEqual((r['category'],r['position'],r['featured']),('coats',7,1))
  with app.connect() as c:
   self.assertEqual(c.execute("SELECT present FROM products WHERE code='MATERIAL'").fetchone()[0],0)
  self.assertNotIn('MATERIAL',[p['code'] for p in self.req(self.admin,'/products?admin=1')['products']])
  self.assertEqual(self.req(customer,'/quotes/'+str(quote['id']))['quote']['total'],quote['total'])
  before=len(self.req(customer,'/quotes')['quotes']);snapshot.write_text(json.dumps({'status':'failed','completeness':'incomplete','rows':rows}))
  with self.assertRaises(ValueError):app.sync(snapshot,config)
  self.assertEqual(len(self.req(customer,'/quotes')['quotes']),before)
 def test_finished_goods_only(self):
  for category in ['毛衣','衣服/成衣刺绣','光板','帽子']:
   self.assertTrue(app.is_finished_category(category),category)
  for category in ['常规辅料','特殊辅料','加加烫画片','布料','衣服/毛衣/辅料','未知分类','']:
   self.assertFalse(app.is_finished_category(category),category)
 def test_admin_category_and_search(self):
  with app.connect() as db:
   original=db.execute("SELECT category FROM products WHERE code='B'").fetchone()[0]
   db.execute("UPDATE products SET category='costumes' WHERE code='B'")
  try:
   d=self.req(self.admin,'/products?admin=1&category=costumes&search=B')
   self.assertEqual([p['code'] for p in d['products']],['B'])
   self.assertEqual(self.req(self.admin,'/products?admin=1&category=costumes&search=nomatch')['total'],0)
   self.assertGreater(self.req(self.admin,'/products?admin=1')['total'],1)
  finally:
   with app.connect() as db:db.execute("UPDATE products SET category=? WHERE code='B'",(original,))
 def test_remote_image_source_restricted(self):
  with app.connect() as c:
   with self.assertRaises(ValueError):app.ensure_image({'photo':'','image_url':'http://127.0.0.1/private','code':'not-real'},c)
 def test_categories_and_feature(self):
  self.req(self.admin,'/admin/products',{'codes':['B'],'changes':{'featured':True,'position':1}});self.assertEqual(self.req(self.client(),'/products?featured=1')['products'][0]['code'],'B')
  cats=self.req(self.admin,'/categories')['categories'];cats[0]['zh']='官网针织';self.req(self.admin,'/admin/categories',{'categories':cats});self.assertEqual(self.req(self.client(),'/categories')['categories'][0]['zh'],'官网针织')
 def test_language_quote_and_downloads(self):
  from pypdf import PdfReader
  import openpyxl
  c,u=self.customer()
  for language,expected in [('zh','zh'),('ru','en'),('es','en')]:
   q=self.create(c,language=language)
   with app.connect() as db:snapshot=json.loads(db.execute('SELECT snapshot FROM quotes WHERE id=?',(q['id'],)).fetchone()[0])
   self.assertEqual(snapshot['language'],expected)
   if language=='zh':
    data=self.req(c,'/quotes/'+str(q['id'])+'/pdf',raw=True)
    self.assertIn('客户报价',''.join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages))
    data=self.req(c,'/quotes/'+str(q['id'])+'/xlsx',raw=True)
    self.assertEqual(openpyxl.load_workbook(io.BytesIO(data),data_only=True).active['C5'].value,'产品名称')
 def test_translated_size_and_color_exports(self):
  customer,_=self.customer()
  with app.connect() as db:
   original=tuple(db.execute("SELECT color,size,carton_qty,carton_size FROM variants WHERE sku='A-M'").fetchone())
   db.execute("UPDATE variants SET color='米白',size='均码',carton_qty=50,carton_size='40 x 30 x 30 cm' WHERE sku='A-M'")
  try:
   dest=app.ROOT/'.local/test-results/language-fix';dest.mkdir(parents=True,exist_ok=True)
   for lang in ['en','zh']:
    q=self.create(customer,[{'sku':'A-M','quantity':2000}],language=lang)
    self.assertEqual(q['lines'][0]['colorEn'],'Off-white');self.assertEqual(q['lines'][0]['size'],'均码')
    for kind in ['pdf','xlsx']:
     data=self.req(customer,'/quotes/'+str(q['id'])+'/'+kind,raw=True);(dest/(lang+'.'+kind)).write_bytes(data)
     if kind=='pdf':
      from pypdf import PdfReader
      text=''.join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
      self.assertIn('均码' if lang=='zh' else 'One size',text)
      if lang=='en':self.assertNotIn('均码',text);self.assertIn('Off-white',text)
      else:self.assertIn('件',text);self.assertIn('昆山加元服装服饰有限公司',text)
     else:
      import openpyxl
      w=openpyxl.load_workbook(io.BytesIO(data),data_only=True)
      self.assertEqual(w.active['E6'].value,'均码' if lang=='zh' else 'One size')
      self.assertEqual(w.active['D6'].value,'米白' if lang=='zh' else 'Off-white')
  finally:
   with app.connect() as db:db.execute("UPDATE variants SET color=?,size=?,carton_qty=?,carton_size=? WHERE sku='A-M'",original)
 def test_publish_blocks_unreviewed_translation(self):
  with app.connect() as db:
   original=db.execute("SELECT color FROM variants WHERE sku='A-M'").fetchone()[0]
   db.execute("UPDATE variants SET color='待核对未知款式' WHERE sku='A-M'")
  try:
   response=self.req(self.admin,'/admin/products',{'codes':['A'],'changes':{'published':True}},400)
   self.assertIn('英文规格',response['error'])
   public=self.req(self.client(),'/products/A')['product']
   self.assertIn('Variant A-M',[v['colorEn'] for v in public['variants']]);self.assertTrue(next(p for p in self.req(self.admin,'/products?admin=1&search=A')['products'] if p['code']=='A')['translationPending'])
  finally:
   with app.connect() as db:db.execute("UPDATE variants SET color=? WHERE sku='A-M'",(original,))
 def test_exports(self):
  c,u=self.customer();q=self.create(c,[{'sku':'A-S','quantity':1900},{'sku':'A-M','quantity':100}]);dest=app.ROOT/'.local/test-results';dest.mkdir(exist_ok=True)
  with app.connect() as db:snapshot=json.loads(db.execute('SELECT snapshot FROM quotes WHERE id=?',(q['id'],)).fetchone()[0])
  (dest/'images').mkdir(exist_ok=True)
  for l in snapshot['lines']:shutil.copy2(app.STATE/'quote-images'/l['image'],dest/'images'/l['image'])
  (dest/'export-input.json').write_text(json.dumps({'quote':snapshot,'imageRoot':str(dest/'images')},ensure_ascii=False))
  for kind in ['pdf','xlsx']:
   data=self.req(c,'/quotes/'+str(q['id'])+'/'+kind,raw=True);(dest/('quote.'+kind)).write_bytes(data)
   if kind=='pdf':
    from pypdf import PdfReader
    reader=PdfReader(io.BytesIO(data));self.assertGreater(len(reader.pages),0);text=''.join(p.extract_text() for p in reader.pages);self.assertIn(q['number'],text);self.assertNotIn('exchangeRate',text)
   else:
    with zipfile.ZipFile(io.BytesIO(data)) as z:self.assertTrue(any(n.startswith('xl/media/') for n in z.namelist()))
    import openpyxl
    w=openpyxl.load_workbook(io.BytesIO(data),data_only=True);self.assertEqual(w.active['H8'].value,q['total'])

 # 本次新增：图片编辑与官网分类均在隔离副本验证。
 def test_product_content_gallery_limits_and_permissions(self):
  import base64
  from PIL import Image
  raw=io.BytesIO();Image.new('RGB',(24,24),'navy').save(raw,'PNG')
  image='data:image/png;base64,'+base64.b64encode(raw.getvalue()).decode()
  self.req(self.client(),'/admin/product-content',{'code':'B','images':[image]},status=403)
  self.req(self.admin,'/admin/product-content',{'code':'B','images':[image]*6},status=400)
  self.req(self.admin,'/admin/product-content',{'code':'B','images':['/etc/passwd']},status=400)
  result=self.req(self.admin,'/admin/product-content',{'code':'B','images':[image,image],'description_zh':'真实产品描述','description_en':'Pet clothing.','guide_en':image})
  self.assertEqual(len(result['content']['images']),1)
  five=[]
  for color in ['red','blue','green','white','black']:
   buf=io.BytesIO();Image.new('RGB',(24,24),color).save(buf,'PNG');five.append('data:image/png;base64,'+base64.b64encode(buf.getvalue()).decode())
  five_result=self.req(self.admin,'/admin/product-content',{'code':'B','images':five})['content']['images'];self.assertEqual(len(five_result),5)
  reordered=list(reversed(five_result))
  self.req(self.admin,'/admin/product-content',{'code':'B','images':reordered})
  self.assertTrue(self.req(self.client(),'/products/B')['product']['image'].endswith(reordered[0]))
  self.req(self.admin,'/admin/product-content',{'code':'B','images':reordered[:1]})
  result={'content':{'images':reordered[:1]}}
  data=self.req(self.client(),'/products/B')['product']
  self.assertEqual(len(data['images']),1);self.assertEqual(data['descriptionEn'],'Pet clothing.')
  self.assertTrue(data['guideEn']);self.assertIsNone(data['guideZh'])
  self.assertTrue(self.req(self.client(),data['images'][0].removeprefix('/api'),raw=True).startswith(b'\xff\xd8'))
  self.req(self.admin,'/admin/product-content',{'code':'B','images':result['content']['images'],'description_en':'中文'},status=400)
  self.req(self.client(),'/products/B/gallery/not-a-product.jpg',status=404)
  # 主图修改不改历史报价图片，保留旧文件。
  self.assertTrue((app.STATE/'images/B.jpg').is_file())

 def test_website_category_and_erp_filter_are_independent(self):
  self.req(self.admin,'/admin/categories',{'categories':[{'id':'test-collection','zh':'测试系列','en':'Test collection'}]})
  self.req(self.admin,'/admin/products',{'codes':['B'],'changes':{'category':'test-collection'}})
  d=self.req(self.admin,'/products?admin=1&sourceCategory='+urllib.parse.quote('毛衣'))
  self.assertIn('毛衣',d['sourceCategories']);self.assertIn('B',[p['code'] for p in d['products']])
  self.assertIn('B',[p['code'] for p in self.req(self.client(),'/products?category=test-collection')['products']])
  self.assertEqual(self.req(self.admin,'/products?admin=1&sourceCategory=test-collection')['total'],0)

 def test_erp_image_pull_deduplicates_and_rejects_other_sources(self):
  from unittest.mock import patch
  with app.connect() as c:
   with patch.object(app,'erp_image_choices',return_value=['https://img.alicdn.com/a.jpg','https://img.alicdn.com/b.jpg']),patch.object(app,'image_bytes',return_value=b'deduplicated-image'):
    result=app.pull_erp_images(c,'C');self.assertEqual(len(result['erpImages']),1)
   with self.assertRaises(ValueError):app.image_bytes('https://example.com/secret',[])

 def test_category_replace_rejects_stale_list(self):
  cats=self.req(self.admin,'/categories')['categories'];ids=[r['id'] for r in cats]
  self.req(self.admin,'/admin/categories',{'categories':cats,'replace':True,'baseIds':ids})
  self.req(self.admin,'/admin/categories',{'categories':cats,'replace':True,'baseIds':['stale']},status=400)
  self.assertEqual([r['id'] for r in self.req(self.admin,'/categories')['categories']],ids)

 def test_admin_publication_filter_and_batch_rollback_reason(self):
  with app.connect() as c:
   old=c.execute("SELECT published FROM products WHERE code='B'").fetchone()[0]
   c.execute("UPDATE products SET published=0 WHERE code='B'")
   c.execute("INSERT INTO products(code,name_zh,name_en,source_category,category) VALUES('FAILIMAGE','缺图测试款','Test style','毛衣','knitwear')")
   c.execute("INSERT INTO variants VALUES('FAILIMAGE-S','FAILIMAGE','红色','S',10,1,100,50,'40x30x30')")
  try:
   unpublished=self.req(self.admin,'/products?admin=1&published=0&sourceCategory='+urllib.parse.quote('毛衣')+'&search=B')
   self.assertEqual([p['code'] for p in unpublished['products']],['B'])
   self.assertEqual(self.req(self.admin,'/products?admin=1&published=1&search=B')['total'],0)
   self.req(self.admin,'/products?admin=1&published=invalid',status=400)
   result=self.req(self.admin,'/admin/products',{'codes':['B','FAILIMAGE'],'changes':{'published':True}},status=400)
   self.assertIn('本批次未上架',result['error']);self.assertIn('FAILIMAGE',result['error']);self.assertIn('缺少产品图片',result['error'])
   with app.connect() as c:self.assertEqual(c.execute("SELECT published FROM products WHERE code='B'").fetchone()[0],0)
   self.assertNotIn('B',[p['code'] for p in self.req(self.client(),'/products?published=0')['products']])
  finally:
   with app.connect() as c:
    c.execute("UPDATE products SET published=? WHERE code='B'",(old,));c.execute("UPDATE products SET present=0 WHERE code='FAILIMAGE'")


 def test_multiple_categories_api_rollup_and_atomic_batch(self):
  from server.categories import category_ids
  with app.connect() as c:original=category_ids(c,'B')
  self.req(self.admin,'/admin/categories',{'categories':[{'id':'multi-root','zh':'测试主分类','en':'Test Parent'},{'id':'multi-child','zh':'测试子分类','en':'Test Child','parentId':'multi-root'},{'id':'multi-other','zh':'测试其他分类','en':'Test Other'}]})
  try:
   self.req(self.admin,'/admin/products',{'codes':['B'],'changes':{'categoryIds':['multi-root','multi-child','multi-other']}})
   for key in ['multi-root','multi-child','multi-other']:
    products=self.req(self.client(),'/products?category='+key)['products']
    self.assertEqual([p['code'] for p in products].count('B'),1)
   self.req(self.admin,'/admin/products',{'codes':['B','missing-style'],'changes':{'categoryIds':['multi-child'],'categoryMode':'remove'}},status=400)
   with app.connect() as c:self.assertIn('multi-child',category_ids(c,'B'))
   self.req(self.admin,'/admin/products',{'codes':['B'],'changes':{'categoryIds':['multi-child'],'categoryMode':'remove'}})
   self.assertNotIn('B',[p['code'] for p in self.req(self.client(),'/products?category=multi-child')['products']])
   self.assertIn('B',[p['code'] for p in self.req(self.client(),'/products?category=multi-other')['products']])
   self.req(self.client(),'/products?admin=1&directCategory=multi-other',status=403)
   self.req(self.admin,'/admin/products',{'codes':['B'],'changes':{'categoryIds':['invalid-category']}},status=400)
  finally:self.req(self.admin,'/admin/products',{'codes':['B'],'changes':{'categoryIds':original}})

 def test_admin_unclassified_filter_and_public_scope(self):
  from server.categories import category_ids
  with app.connect() as c:original=category_ids(c,'B')
  try:
   self.req(self.admin,'/admin/products',{'codes':['B'],'changes':{'categoryIds':[]}})
   self.assertEqual([p['code'] for p in self.req(self.admin,'/products?admin=1&unclassified=1&search=B')['products']],['B'])
   self.req(self.admin,'/admin/products',{'codes':['B'],'changes':{'categoryIds':['knitwear']}})
   self.assertEqual(self.req(self.admin,'/products?admin=1&unclassified=1&search=B')['total'],0)
   self.assertIn('B',[p['code'] for p in self.req(self.client(),'/products?unclassified=1&search=B')['products']])
  finally:self.req(self.admin,'/admin/products',{'codes':['B'],'changes':{'categoryIds':original}})

if __name__=='__main__':unittest.main(verbosity=2)
